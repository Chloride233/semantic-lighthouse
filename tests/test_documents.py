from pathlib import Path
from datetime import timedelta
from hashlib import sha256
from io import BytesIO

from fastapi.testclient import TestClient
from docx import Document as DocxDocument
from sqlalchemy import select
from sqlalchemy.orm import Session

from conftest import register_and_login
from semantic_lighthouse.config import Settings, get_settings
from semantic_lighthouse.models import DocumentUploadSession, utc_now


def _create_group(client: TestClient, headers: dict[str, str], name: str = "Team") -> str:
    response = client.post("/groups", json={"name": name}, headers=headers)
    assert response.status_code == 201
    return response.json()["id"]


def _join_group(client: TestClient, group_id: str, owner_headers: dict[str, str], member_headers: dict[str, str]) -> None:
    invite = client.post(f"/groups/{group_id}/invites", headers=owner_headers)
    assert invite.status_code == 201
    joined = client.post(
        "/groups/join-by-invite",
        json={"invite_code": invite.json()["invite_code"]},
        headers=member_headers,
    )
    assert joined.status_code == 200


def _override_settings(client: TestClient, knowledge_base_path: Path, max_upload_bytes: int = 1024 * 1024) -> None:
    def override() -> Settings:
        return Settings(
            database_url="sqlite+pysqlite:///:memory:",
            jwt_secret_key="test-secret-key-at-least-32-bytes",
            cookie_secure=False,
            cookie_samesite="lax",
            knowledge_base_path=str(knowledge_base_path),
            max_markdown_upload_bytes=max_upload_bytes,
            document_storage_path=str(knowledge_base_path / "document-storage"),
            upload_tmp_path=str(knowledge_base_path / "upload-tmp"),
            max_document_upload_bytes=50 * 1024 * 1024,
            upload_chunk_bytes=8,
            embedding_provider="fake",
        )

    client.app.dependency_overrides[get_settings] = override


def _write_markdown(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")


def test_owner_imports_local_knowledge_base_and_duplicate_import_is_idempotent(client, tmp_path):
    _, _, owner_headers = register_and_login(client, "owner@example.com")
    group_id = _create_group(client, owner_headers)
    knowledge_base = tmp_path / "knowledge-graph"
    _write_markdown(
        knowledge_base / "concepts" / "ontology.md",
        """---
title: Ontology
entityType: Concept
source: official-doc
status: reviewed
---

# Ontology

Ontology connects business objects and operations.
""",
    )
    _write_markdown(knowledge_base / "INDEX.md", "# Index")
    _write_markdown(knowledge_base / "schema.md", "# Schema")
    _write_markdown(knowledge_base / "concepts" / "_TEMPLATE.md", "# Template")
    _override_settings(client, knowledge_base)

    response = client.post(f"/groups/{group_id}/documents/import-local", headers=owner_headers)

    assert response.status_code == 201
    payload = response.json()
    assert payload["imported_count"] == 1
    assert payload["skipped_count"] == 3
    assert payload["documents"][0]["title"] == "Ontology"
    assert payload["documents"][0]["source_path"] == "concepts/ontology.md"
    assert payload["documents"][0]["frontmatter"]["entityType"] == "Concept"

    duplicate = client.post(f"/groups/{group_id}/documents/import-local", headers=owner_headers)

    assert duplicate.status_code == 201
    assert duplicate.json()["imported_count"] == 0
    documents = client.get(f"/groups/{group_id}/documents", headers=owner_headers)
    assert len(documents.json()) == 1


def test_member_cannot_import_or_upload_documents(client, tmp_path):
    _, _, owner_headers = register_and_login(client, "owner@example.com")
    _, _, member_headers = register_and_login(client, "member@example.com")
    group_id = _create_group(client, owner_headers)
    _join_group(client, group_id, owner_headers, member_headers)
    _override_settings(client, tmp_path)

    import_response = client.post(f"/groups/{group_id}/documents/import-local", headers=member_headers)
    upload_response = client.post(
        f"/groups/{group_id}/documents/upload",
        files={"file": ("note.md", b"# Note\n\ncontent", "text/markdown")},
        headers=member_headers,
    )

    assert import_response.status_code == 403
    assert upload_response.status_code == 403


def test_upload_rejects_non_markdown_and_large_markdown(client, tmp_path):
    _, _, owner_headers = register_and_login(client, "owner@example.com")
    group_id = _create_group(client, owner_headers)
    _override_settings(client, tmp_path, max_upload_bytes=10)

    non_markdown = client.post(
        f"/groups/{group_id}/documents/upload",
        files={"file": ("note.exe", b"plain text", "application/octet-stream")},
        headers=owner_headers,
    )
    too_large = client.post(
        f"/groups/{group_id}/documents/upload",
        files={"file": ("note.md", b"# Title\n\nThis is too large.", "text/markdown")},
        headers=owner_headers,
    )

    assert non_markdown.status_code == 400
    assert too_large.status_code == 413


def test_upload_supports_txt_docx_and_pdf(client, tmp_path):
    _, _, owner_headers = register_and_login(client, "owner@example.com")
    group_id = _create_group(client, owner_headers)
    _override_settings(client, tmp_path)

    txt = client.post(
        f"/groups/{group_id}/documents/upload",
        files={"file": ("plain.txt", b"plain ontology text", "text/plain")},
        headers=owner_headers,
    )
    docx = client.post(
        f"/groups/{group_id}/documents/upload",
        files={"file": ("brief.docx", _docx_bytes("docx ontology text"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=owner_headers,
    )
    pdf = client.post(
        f"/groups/{group_id}/documents/upload",
        files={"file": ("paper.pdf", _pdf_bytes(), "application/pdf")},
        headers=owner_headers,
    )

    assert txt.status_code == 201
    assert txt.json()["parser"] == "text"
    assert docx.status_code == 201
    assert docx.json()["parser"] == "python-docx"
    assert pdf.status_code == 201
    assert pdf.json()["parser"] == "pypdf"

    search = client.get(f"/groups/{group_id}/documents/search", params={"q": "ontology"}, headers=owner_headers)
    assert search.status_code == 200
    assert {result["title"] for result in search.json()} >= {"plain", "brief", "paper"}


def test_upload_parses_frontmatter_headings_and_search_returns_citation_fields(client, tmp_path):
    _, _, owner_headers = register_and_login(client, "owner@example.com")
    _, _, member_headers = register_and_login(client, "member@example.com")
    group_id = _create_group(client, owner_headers)
    _join_group(client, group_id, owner_headers, member_headers)
    _override_settings(client, tmp_path)
    markdown = b"""---
title: AI Roadmap
entityType: Methodology
source: personal-analysis
status: draft
---

# AI Transformation

Data foundation supports ontology and retrieval.

## Risk

Permission boundaries protect citations.
"""

    upload = client.post(
        f"/groups/{group_id}/documents/upload",
        files={"file": ("roadmap.md", markdown, "text/markdown")},
        headers=owner_headers,
    )
    document_id = upload.json()["id"]
    detail = client.get(f"/groups/{group_id}/documents/{document_id}", headers=member_headers)
    search = client.get(f"/groups/{group_id}/documents/search", params={"q": "ontology"}, headers=member_headers)

    assert upload.status_code == 201
    assert upload.json()["frontmatter"]["entityType"] == "Methodology"
    assert detail.status_code == 200
    assert detail.json()["chunks"][0]["heading_path"] == "AI Transformation"
    assert search.status_code == 200
    result = search.json()[0]
    assert result["document_id"] == document_id
    assert result["chunk_id"]
    assert result["title"] == "AI Roadmap"
    assert result["source_path"] == "upload:roadmap.md"
    assert result["entity_type"] == "Methodology"
    assert "ontology" in result["snippet"]


def test_non_member_cannot_access_documents_or_search(client, tmp_path):
    _, _, owner_headers = register_and_login(client, "owner@example.com")
    _, _, outsider_headers = register_and_login(client, "outsider@example.com")
    group_id = _create_group(client, owner_headers)
    _override_settings(client, tmp_path)

    list_response = client.get(f"/groups/{group_id}/documents", headers=outsider_headers)
    search_response = client.get(
        f"/groups/{group_id}/documents/search",
        params={"q": "anything"},
        headers=outsider_headers,
    )

    assert list_response.status_code == 403
    assert search_response.status_code == 403


def test_search_is_filtered_by_group_id(client, tmp_path):
    _, _, owner_headers = register_and_login(client, "owner@example.com")
    group_a = _create_group(client, owner_headers, "A")
    group_b = _create_group(client, owner_headers, "B")
    _override_settings(client, tmp_path)

    client.post(
        f"/groups/{group_a}/documents/upload",
        files={"file": ("a.md", b"# A\n\nshared keyword belongs to group A", "text/markdown")},
        headers=owner_headers,
    )
    client.post(
        f"/groups/{group_b}/documents/upload",
        files={"file": ("b.md", b"# B\n\nshared keyword belongs to group B", "text/markdown")},
        headers=owner_headers,
    )

    response = client.get(f"/groups/{group_a}/documents/search", params={"q": "shared"}, headers=owner_headers)

    assert response.status_code == 200
    assert len(response.json()) == 1
    assert response.json()[0]["source_path"] == "upload:a.md"


def test_chunked_upload_session_complete_and_instant_upload(client, tmp_path):
    _, _, owner_headers = register_and_login(client, "owner@example.com")
    group_id = _create_group(client, owner_headers)
    _override_settings(client, tmp_path)
    content = b"# Ontology\n\nchunked ontology upload works"
    file_hash = sha256(content).hexdigest()

    init = client.post(
        f"/groups/{group_id}/documents/uploads/init",
        json={"file_name": "chunked.md", "file_size": len(content), "file_hash": file_hash, "chunk_size": 16},
        headers=owner_headers,
    )
    upload_id = init.json()["upload_id"]
    first = client.put(
        f"/groups/{group_id}/documents/uploads/{upload_id}/chunks/0",
        files={"file": ("0.part", content[:16], "application/octet-stream")},
        headers=owner_headers,
    )
    second = client.put(
        f"/groups/{group_id}/documents/uploads/{upload_id}/chunks/1",
        files={"file": ("1.part", content[16:32], "application/octet-stream")},
        headers=owner_headers,
    )
    third = client.put(
        f"/groups/{group_id}/documents/uploads/{upload_id}/chunks/2",
        files={"file": ("2.part", content[32:], "application/octet-stream")},
        headers=owner_headers,
    )
    complete = client.post(f"/groups/{group_id}/documents/uploads/{upload_id}/complete", headers=owner_headers)
    instant = client.post(
        f"/groups/{group_id}/documents/uploads/init",
        json={"file_name": "chunked.md", "file_size": len(content), "file_hash": file_hash, "chunk_size": 16},
        headers=owner_headers,
    )

    assert init.status_code == 201
    assert init.json()["type"] == "UPLOAD_SESSION"
    assert init.json()["total_chunks"] == 3
    assert first.json()["uploaded_count"] == 1
    assert second.json()["uploaded_count"] == 2
    assert third.json()["uploaded_count"] == 3
    assert complete.status_code == 200
    assert complete.json()["title"] == "Ontology"
    assert complete.json()["file_hash"] == file_hash
    assert instant.status_code == 201
    assert instant.json()["type"] == "INSTANT"
    assert instant.json()["document_id"] == complete.json()["id"]


def test_chunked_upload_resume_and_idempotent_chunk(client, tmp_path):
    _, _, owner_headers = register_and_login(client, "owner@example.com")
    group_id = _create_group(client, owner_headers)
    _override_settings(client, tmp_path)
    content = b"# Resume\n\nresumable upload"
    file_hash = sha256(content).hexdigest()

    init = client.post(
        f"/groups/{group_id}/documents/uploads/init",
        json={"file_name": "resume.md", "file_size": len(content), "file_hash": file_hash, "chunk_size": 8},
        headers=owner_headers,
    )
    upload_id = init.json()["upload_id"]
    first = client.put(
        f"/groups/{group_id}/documents/uploads/{upload_id}/chunks/0",
        files={"file": ("0.part", content[:8], "application/octet-stream")},
        headers=owner_headers,
    )
    repeated = client.put(
        f"/groups/{group_id}/documents/uploads/{upload_id}/chunks/0",
        files={"file": ("0.part", content[:8], "application/octet-stream")},
        headers=owner_headers,
    )
    resumed = client.post(
        f"/groups/{group_id}/documents/uploads/init",
        json={"file_name": "resume.md", "file_size": len(content), "file_hash": file_hash, "chunk_size": 8},
        headers=owner_headers,
    )

    assert first.status_code == 200
    assert repeated.status_code == 200
    assert repeated.json()["uploaded_count"] == 1
    assert resumed.status_code == 201
    assert resumed.json()["type"] == "UPLOAD_SESSION"
    assert resumed.json()["upload_id"] == upload_id
    assert resumed.json()["uploaded_chunks"] == [0]


def test_chunked_upload_rejects_incomplete_hash_mismatch_and_expired_session(client, db_session: Session, tmp_path):
    _, _, owner_headers = register_and_login(client, "owner@example.com")
    group_id = _create_group(client, owner_headers)
    _override_settings(client, tmp_path)
    content = b"# Bad\n\nhash mismatch"
    wrong_hash = sha256(b"wrong").hexdigest()

    init = client.post(
        f"/groups/{group_id}/documents/uploads/init",
        json={"file_name": "bad.md", "file_size": len(content), "file_hash": wrong_hash, "chunk_size": 8},
        headers=owner_headers,
    )
    upload_id = init.json()["upload_id"]
    incomplete = client.post(f"/groups/{group_id}/documents/uploads/{upload_id}/complete", headers=owner_headers)
    client.put(
        f"/groups/{group_id}/documents/uploads/{upload_id}/chunks/0",
        files={"file": ("0.part", content[:8], "application/octet-stream")},
        headers=owner_headers,
    )
    client.put(
        f"/groups/{group_id}/documents/uploads/{upload_id}/chunks/1",
        files={"file": ("1.part", content[8:16], "application/octet-stream")},
        headers=owner_headers,
    )
    client.put(
        f"/groups/{group_id}/documents/uploads/{upload_id}/chunks/2",
        files={"file": ("2.part", content[16:], "application/octet-stream")},
        headers=owner_headers,
    )
    mismatch = client.post(f"/groups/{group_id}/documents/uploads/{upload_id}/complete", headers=owner_headers)
    status_response = client.get(f"/groups/{group_id}/documents/uploads/{upload_id}", headers=owner_headers)

    expired_init = client.post(
        f"/groups/{group_id}/documents/uploads/init",
        json={"file_name": "expired.md", "file_size": 10, "file_hash": sha256(b"expired!!").hexdigest(), "chunk_size": 5},
        headers=owner_headers,
    )
    expired_id = expired_init.json()["upload_id"]
    expired_session = db_session.scalar(select(DocumentUploadSession).where(DocumentUploadSession.id == expired_id))
    assert expired_session is not None
    expired_session.expires_at = utc_now() - timedelta(hours=1)
    db_session.commit()
    expired_upload = client.put(
        f"/groups/{group_id}/documents/uploads/{expired_id}/chunks/0",
        files={"file": ("0.part", b"12345", "application/octet-stream")},
        headers=owner_headers,
    )

    assert incomplete.status_code == 400
    assert mismatch.status_code == 400
    assert status_response.json()["status"] == "failed"
    assert expired_upload.status_code == 410


def test_member_cannot_use_chunked_upload(client, tmp_path):
    _, _, owner_headers = register_and_login(client, "owner@example.com")
    _, _, member_headers = register_and_login(client, "member@example.com")
    group_id = _create_group(client, owner_headers)
    _join_group(client, group_id, owner_headers, member_headers)
    _override_settings(client, tmp_path)
    content = b"# No\n\nmember upload"

    response = client.post(
        f"/groups/{group_id}/documents/uploads/init",
        json={"file_name": "member.md", "file_size": len(content), "file_hash": sha256(content).hexdigest()},
        headers=member_headers,
    )

    assert response.status_code == 403


def test_upload_temp_chunks_cleaned_after_successful_complete(client, tmp_path):
    """Temp chunk files must be removed after complete; merged file is kept."""
    _, _, owner_headers = register_and_login(client, "owner@example.com")
    group_id = _create_group(client, owner_headers)
    _override_settings(client, tmp_path)
    content = b"# Cleanup\n\nsuccess path"
    file_hash = sha256(content).hexdigest()

    init = client.post(
        f"/groups/{group_id}/documents/uploads/init",
        json={"file_name": "cleanup.md", "file_size": len(content), "file_hash": file_hash, "chunk_size": 8},
        headers=owner_headers,
    )
    upload_id = init.json()["upload_id"]
    for start in range(0, len(content), 8):
        end = min(start + 8, len(content))
        client.put(
            f"/groups/{group_id}/documents/uploads/{upload_id}/chunks/{start // 8}",
            files={"file": (f"{start // 8}.part", content[start:end], "application/octet-stream")},
            headers=owner_headers,
        )
    complete = client.post(f"/groups/{group_id}/documents/uploads/{upload_id}/complete", headers=owner_headers)
    assert complete.status_code == 200

    temp_dir = tmp_path / "upload-tmp" / group_id / upload_id
    assert not temp_dir.exists(), f"Temp chunk dir should be cleaned: {temp_dir}"

    doc_storage_dir = tmp_path / "document-storage" / group_id / upload_id
    assert doc_storage_dir.exists(), "Merged file in document-storage should be kept after success"


def test_upload_temp_and_merged_cleaned_after_hash_mismatch(client, tmp_path):
    """Both temp chunks and the merged file must be removed after hash mismatch."""
    _, _, owner_headers = register_and_login(client, "owner@example.com")
    group_id = _create_group(client, owner_headers)
    _override_settings(client, tmp_path)
    content = b"# Mismatch\n\nhash failure path"
    wrong_hash = sha256(b"wrong").hexdigest()

    init = client.post(
        f"/groups/{group_id}/documents/uploads/init",
        json={"file_name": "mismatch.md", "file_size": len(content), "file_hash": wrong_hash, "chunk_size": 8},
        headers=owner_headers,
    )
    upload_id = init.json()["upload_id"]
    for start in range(0, len(content), 8):
        end = min(start + 8, len(content))
        client.put(
            f"/groups/{group_id}/documents/uploads/{upload_id}/chunks/{start // 8}",
            files={"file": (f"{start // 8}.part", content[start:end], "application/octet-stream")},
            headers=owner_headers,
        )
    mismatch = client.post(f"/groups/{group_id}/documents/uploads/{upload_id}/complete", headers=owner_headers)
    assert mismatch.status_code == 400

    temp_dir = tmp_path / "upload-tmp" / group_id / upload_id
    assert not temp_dir.exists(), f"Temp chunk dir should be cleaned after hash mismatch: {temp_dir}"

    doc_storage_dir = tmp_path / "document-storage" / group_id / upload_id
    assert not doc_storage_dir.exists(), f"Merged file should be cleaned after hash mismatch: {doc_storage_dir}"


def test_upload_temp_and_merged_cleaned_after_parser_failure(client, tmp_path):
    """Both temp chunks and the merged file must be removed after parse failure."""
    _, _, owner_headers = register_and_login(client, "owner@example.com")
    group_id = _create_group(client, owner_headers)
    _override_settings(client, tmp_path)
    content = _empty_docx_bytes()
    file_hash = sha256(content).hexdigest()

    init = client.post(
        f"/groups/{group_id}/documents/uploads/init",
        json={"file_name": "empty.docx", "file_size": len(content), "file_hash": file_hash, "chunk_size": 4096},
        headers=owner_headers,
    )
    upload_id = init.json()["upload_id"]
    client.put(
        f"/groups/{group_id}/documents/uploads/{upload_id}/chunks/0",
        files={"file": ("0.part", content, "application/octet-stream")},
        headers=owner_headers,
    )
    parse_fail = client.post(f"/groups/{group_id}/documents/uploads/{upload_id}/complete", headers=owner_headers)
    assert parse_fail.status_code == 400

    temp_dir = tmp_path / "upload-tmp" / group_id / upload_id
    assert not temp_dir.exists(), f"Temp chunk dir should be cleaned after parse failure: {temp_dir}"

    doc_storage_dir = tmp_path / "document-storage" / group_id / upload_id
    assert not doc_storage_dir.exists(), f"Merged file should be cleaned after parse failure: {doc_storage_dir}"


def test_member_cannot_get_upload_session(client, tmp_path):
    """GET upload session requires Owner/Admin; Member access is denied."""
    _, _, owner_headers = register_and_login(client, "owner@example.com")
    _, _, member_headers = register_and_login(client, "member@example.com")
    group_id = _create_group(client, owner_headers)
    _join_group(client, group_id, owner_headers, member_headers)
    _override_settings(client, tmp_path)
    content = b"# Session\n\nvisibility test"

    init = client.post(
        f"/groups/{group_id}/documents/uploads/init",
        json={"file_name": "session.md", "file_size": len(content), "file_hash": sha256(content).hexdigest()},
        headers=owner_headers,
    )
    upload_id = init.json()["upload_id"]

    member_get = client.get(f"/groups/{group_id}/documents/uploads/{upload_id}", headers=member_headers)
    assert member_get.status_code == 403, "Member should not be able to view upload session"


def test_recomplete_returns_same_document(client, tmp_path):
    """Calling complete twice on the same upload returns the same document."""
    _, _, owner_headers = register_and_login(client, "owner@example.com")
    group_id = _create_group(client, owner_headers)
    _override_settings(client, tmp_path)
    content = b"# Recomplete\n\nidempotent complete"
    file_hash = sha256(content).hexdigest()

    init = client.post(
        f"/groups/{group_id}/documents/uploads/init",
        json={"file_name": "recomplete.md", "file_size": len(content), "file_hash": file_hash, "chunk_size": 8},
        headers=owner_headers,
    )
    upload_id = init.json()["upload_id"]
    for start in range(0, len(content), 8):
        end = min(start + 8, len(content))
        client.put(
            f"/groups/{group_id}/documents/uploads/{upload_id}/chunks/{start // 8}",
            files={"file": (f"{start // 8}.part", content[start:end], "application/octet-stream")},
            headers=owner_headers,
        )
    first = client.post(f"/groups/{group_id}/documents/uploads/{upload_id}/complete", headers=owner_headers)
    assert first.status_code == 200
    first_doc_id = first.json()["id"]

    second = client.post(f"/groups/{group_id}/documents/uploads/{upload_id}/complete", headers=owner_headers)
    assert second.status_code == 200
    assert second.json()["id"] == first_doc_id, "Re-complete should return the same document"


def test_chunk_rejected_when_session_not_uploading(client, db_session: Session, tmp_path):
    """Chunk upload and complete must be rejected when session is not in 'uploading' status."""
    _, _, owner_headers = register_and_login(client, "owner@example.com")
    group_id = _create_group(client, owner_headers)
    _override_settings(client, tmp_path)
    content = b"# Stale\n\nsession state"
    file_hash = sha256(content).hexdigest()

    init = client.post(
        f"/groups/{group_id}/documents/uploads/init",
        json={"file_name": "stale.md", "file_size": len(content), "file_hash": file_hash, "chunk_size": 8},
        headers=owner_headers,
    )
    upload_id = init.json()["upload_id"]

    # Manually set session to "completing" to simulate a concurrent complete.
    upload_session = db_session.scalar(select(DocumentUploadSession).where(DocumentUploadSession.id == upload_id))
    assert upload_session is not None
    upload_session.status = "completing"
    db_session.commit()

    chunk_rejected = client.put(
        f"/groups/{group_id}/documents/uploads/{upload_id}/chunks/0",
        files={"file": ("0.part", content[:8], "application/octet-stream")},
        headers=owner_headers,
    )
    complete_rejected = client.post(f"/groups/{group_id}/documents/uploads/{upload_id}/complete", headers=owner_headers)

    assert chunk_rejected.status_code == 409, "Chunk upload should be rejected when session is completing"
    assert complete_rejected.status_code == 409, "Complete should be rejected when session is completing"


def _empty_docx_bytes() -> bytes:
    """Return a valid DOCX file with no text paragraphs (triggers parse failure)."""
    buffer = BytesIO()
    document = DocxDocument()
    document.save(buffer)
    return buffer.getvalue()


def _docx_bytes(text: str) -> bytes:
    buffer = BytesIO()
    document = DocxDocument()
    document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


def _pdf_bytes() -> bytes:
    return b"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>
endobj
4 0 obj
<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>
endobj
5 0 obj
<< /Length 64 >>
stream
BT
/F1 24 Tf
100 700 Td
(PDF ontology extraction works) Tj
ET
endstream
endobj
xref
0 6
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000241 00000 n 
0000000311 00000 n 
trailer
<< /Size 6 /Root 1 0 R >>
startxref
432
%%EOF
"""
