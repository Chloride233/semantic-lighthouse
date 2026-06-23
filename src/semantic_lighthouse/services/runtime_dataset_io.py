"""Runtime dataset IO helpers — path validation and CSV/XLSX row readers.

Extracted from services/runtime.py per R1C.
No API, permission, migration, or query semantic changes.
"""

from __future__ import annotations

import csv
from pathlib import Path

from semantic_lighthouse.config import get_settings


# ═══════════════════════════════════════════════════════════════════════════
#  path safety
# ═══════════════════════════════════════════════════════════════════════════


def _validate_dataset_path(
    storage_path: str, group_id: str, project_id: str,
) -> Path:
    """Resolve storage_path and verify it stays within the dataset storage tree."""
    settings = get_settings()
    root = Path(settings.dataset_storage_path).resolve()
    target = Path(storage_path).resolve()

    try:
        target.relative_to(root)
    except ValueError:
        raise ValueError("Dataset path is outside the storage root")

    expected_prefix = root / group_id / project_id
    try:
        target.relative_to(expected_prefix)
    except ValueError:
        raise ValueError("Dataset path is outside the project scope")

    return target


# ═══════════════════════════════════════════════════════════════════════════
#  CSV / XLSX row reader (streaming for CSV, read_only for XLSX)
# ═══════════════════════════════════════════════════════════════════════════


def _stream_csv_rows(
    file_path: str,
    max_rows: int,
) -> tuple[list[str], list[list[str | None]], int, bool]:
    """Stream rows from a CSV file. Never loads entire file into memory.

    Returns (header, data_rows, scanned_count, truncated).
    """
    scanned = 0
    rows: list[list[str | None]] = []

    with open(file_path, "r", encoding="utf-8-sig", errors="replace") as fh:
        reader = csv.reader(fh, strict=True)
        try:
            header_row = next(reader)
        except StopIteration:
            raise ValueError("CSV file has no rows")
        except csv.Error as e:
            raise ValueError(f"CSV parse error: {e}") from e

        header = [h.strip() for h in header_row]
        if not header or all(h == "" for h in header):
            raise ValueError("CSV file has no valid header row")

        # Deduplicate check
        seen: set[str] = set()
        for h in header:
            if not h:
                raise ValueError("CSV header contains empty column name")
            if h in seen:
                raise ValueError(f"Duplicate column name: {h!r}")
            seen.add(h)

        for row in reader:
            scanned += 1
            if len(rows) >= max_rows:
                return header, rows, scanned, True

            parsed: list[str | None] = []
            for i in range(len(header)):
                if i >= len(row) or row[i].strip() == "":
                    parsed.append(None)
                else:
                    parsed.append(row[i].strip())
            rows.append(parsed)

    return header, rows, scanned, scanned > 0 and len(rows) >= max_rows


def _stream_xlsx_rows(
    file_path: str,
    max_rows: int,
) -> tuple[list[str], list[list[str | None]], int, bool]:
    """Read rows from an XLSX file (first visible sheet, read_only)."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportError("openpyxl is required for XLSX reading") from None

    try:
        wb = load_workbook(file_path, read_only=True, data_only=True, keep_links=False)
    except Exception as e:
        raise ValueError(f"Cannot open XLSX file: {e}") from e

    try:
        ws = None
        for name in wb.sheetnames:
            candidate = wb[name]
            if candidate and candidate.sheet_state == "visible":
                ws = candidate
                break
        if ws is None:
            raise ValueError("XLSX file has no visible worksheets")

        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if header_row is None or all(v is None or str(v).strip() == "" for v in header_row):
            raise ValueError("XLSX file has no valid header row")

        header = [str(v).strip() if v is not None else "" for v in header_row]
        while header and header[-1] == "":
            header.pop()
        if not header:
            raise ValueError("XLSX file has no valid header row")

        scanned = 0
        rows: list[list[str | None]] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            scanned += 1
            if len(rows) >= max_rows:
                return header, rows, scanned, True

            parsed: list[str | None] = []
            for i in range(len(header)):
                if i >= len(row) or row[i] is None or str(row[i]).strip() == "":
                    parsed.append(None)
                else:
                    parsed.append(str(row[i]).strip())
            rows.append(parsed)

        return header, rows, scanned, False
    finally:
        wb.close()


def _read_dataset_rows(
    file_path: str,
    file_format: str,
    max_rows: int,
) -> tuple[list[str], list[list[str | None]], int, bool]:
    """Read rows from a dataset file.

    Returns (header, data_rows, scanned_count, truncated).
    CSV uses streaming; XLSX uses openpyxl read_only.
    """
    if file_format == "csv":
        return _stream_csv_rows(file_path, max_rows)
    if file_format == "xlsx":
        return _stream_xlsx_rows(file_path, max_rows)
    raise ValueError(f"Unsupported file format: {file_format}")
